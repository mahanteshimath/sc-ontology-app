"""One-shot fill of docs/3M-ONTOLOGIST/3M-ONTOLOGIST.pptx from its own blank template.

Reuses the template's own background images, fonts (Manrope) and text color (#202729) so
the generated deck stays visually identical to the theme the template shipped with.
Not idempotent: it removes a slide and adds text boxes, so re-running it against its own
output will duplicate content. Restore the template from git before re-running.
"""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE

PATH = r"docs\3M-ONTOLOGIST\3M-ONTOLOGIST.pptx"
TEXT_COLOR = RGBColor(0x20, 0x27, 0x29)
FONT = "Manrope"
BOX_FILL = RGBColor(0x42, 0x85, 0xF4)   # theme accent1
BOX_FILL_DARK = RGBColor(0x21, 0x21, 0x21)  # theme accent2

TEAM_NAME = "3M-Ontologist"
TEAM_LEADER = "Mahantesh Hiremath"
TEAM_SIZE = "3"
PROBLEM_STATEMENT = "Governed conversational analytics for supply chain — one metric definition, provably identical everywhere it's asked."


def set_label_value(shape, value):
    """Append ' <value>' to a 'Label :' textbox as a second run, same font as the label."""
    para = shape.text_frame.paragraphs[0]
    label_run = para.runs[0]
    run = para.add_run()
    run.text = " " + value
    run.font.size = label_run.font.size
    run.font.bold = False
    run.font.name = FONT
    run.font.color.rgb = TEXT_COLOR


def add_content_textbox(slide, left, top, width, height, blocks):
    """blocks: list of (text, bold) paragraphs into one textbox, matching the guidelines-slide style."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, (text, bold) in enumerate(blocks):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = para.add_run()
        run.text = text
        run.font.size = Pt(15)
        run.font.bold = bold
        run.font.name = FONT
        run.font.color.rgb = TEXT_COLOR
    return box


def add_box(slide, left, top, width, height, text, fill, font_color=RGBColor(0xFF, 0xFF, 0xFF), size=11):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    tf = shape.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.name = FONT
    run.font.color.rgb = font_color
    return shape


def add_arrow(slide, left, top, width, height, rotation=0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    shape.rotation = rotation
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x59, 0x59, 0x59)
    shape.line.fill.background()
    return shape


def remove_slide(prs, index):
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[index])


def build():
    prs = Presentation(PATH)

    # --- Slide 0: title ---
    title_slide = prs.slides[0]
    for shape in title_slide.shapes:
        if not shape.has_text_frame:
            continue
        label = shape.text_frame.text
        if label.startswith("Team Name"):
            set_label_value(shape, TEAM_NAME)
        elif label.startswith("Problem Statement"):
            set_label_value(shape, PROBLEM_STATEMENT)
        elif label.startswith("Team Leader Name"):
            set_label_value(shape, TEAM_LEADER)
        elif label.startswith("Team Size"):
            set_label_value(shape, TEAM_SIZE)

    # --- Remove the "Submission Guidelines" instructions slide (was index 1) ---
    remove_slide(prs, 1)

    # After removal: index1=Problem Brief, index2=Architecture, index3=Additional/Impact, index4=Closing
    problem_slide = prs.slides[1]
    arch_slide = prs.slides[2]
    impact_slide = prs.slides[3]

    # --- Problem Brief ---
    add_content_textbox(problem_slide, Emu(329100), Emu(400000), Emu(8485800), Emu(400000), [
        ("Problem Brief", True),
    ])
    add_content_textbox(problem_slide, Emu(329100), Emu(950000), Emu(8485800), Emu(3900000), [
        ("What real business problem does this solve?", True),
        ("Supply-chain metrics disagree across ERP, logistics, supplier and IoT systems — the same "
         "\u201con-time delivery\u201d returns different numbers to different teams, and the gap is usually "
         "discovered in a meeting instead of a test.", False),
        ("Who is the target user / persona?", True),
        ("Supply chain planners, procurement analysts and logistics leads who make S&OP and "
         "supplier-scorecard decisions off a shared metric.", False),
        ("What is the current pain point, and how does this improve it?", True),
        ("A 0.0068 metric spread silently clears 12 of 300 suppliers that shouldn't pass \u2014 22% of "
         "the \u201ccompliant\u201d list \u2014 and never over-escalates, so nobody complains and the error survives "
         "indefinitely. We define every metric once, bind it to every consumer, and drift-test each "
         "binding against an independent canonical calculation on a schedule.", False),
        ("Industry / domain context", True),
        ("Manufacturing supply chain \u2014 ERP purchase and sales orders, logistics/carrier systems, "
         "supplier systems, and IoT shipment telemetry, at multi-million-row scale.", False),
    ])

    # --- Architecture Diagram ---
    add_content_textbox(arch_slide, Emu(329100), Emu(300000), Emu(8485800), Emu(400000), [
        ("Architecture \u2014 one governed execution path", True),
    ])
    box_w, box_h, gap = Emu(1650000), Emu(850000), Emu(230000)
    box_y = Emu(900000)
    x = Emu(329100)
    labels = ["RAW\nERP \u00b7 logistics \u00b7 supplier \u00b7 IoT", "CANONICAL\natomic-grain facts",
              "SEMANTIC\nsingle query surface", "GOVERNANCE\nregistry + drift test + eval"]
    box_positions = []
    for i, label in enumerate(labels):
        add_box(arch_slide, x, box_y, box_w, box_h, label, BOX_FILL, size=10)
        box_positions.append(x)
        if i < len(labels) - 1:
            add_arrow(arch_slide, Emu(x + box_w), Emu(box_y + box_h // 2 - 100000), gap, Emu(200000))
        x = Emu(x + box_w + gap)

    consumer_y = Emu(box_y + box_h + Emu(500000))
    add_arrow(arch_slide, Emu(box_positions[2] + box_w // 2 - 100000), Emu(box_y + box_h), Emu(200000), Emu(300000), rotation=90)
    add_arrow(arch_slide, Emu(box_positions[3] + box_w // 2 - 100000), Emu(box_y + box_h), Emu(200000), Emu(300000), rotation=90)
    add_box(arch_slide, Emu(box_positions[2] - 300000), consumer_y, Emu(1900000), Emu(650000), "Next.js App\n(dashboards + chat)", BOX_FILL_DARK, size=9)
    add_box(arch_slide, Emu(box_positions[3] - 300000), consumer_y, Emu(1900000), Emu(650000), "Cortex Agent\n(verified-query bound)", BOX_FILL_DARK, size=9)

    add_content_textbox(arch_slide, Emu(329100), Emu(3800000), Emu(8485800), Emu(1250000), [
        ("Data sources: ERP (POs, sales orders), logistics/carrier systems, supplier systems, IoT "
         "shipment telemetry \u2014 structured and semi-structured.", False),
        ("Modular components: metric registry, drift-test engine, persona-scoped Snowflake roles + "
         "row access policies, Cortex Agent, Next.js app \u2014 each independently swappable.", False),
        ("The load-bearing arrow: the drift test compares every semantic view against an independent "
         "canonical SQL calculation on the atomic fact \u2014 not just cross-view agreement.", False),
    ])

    # --- Impact Statement (repurposes the "Additional Slide" placeholder, moved to a header position) ---
    for shape in impact_slide.shapes:
        if shape.has_text_frame and shape.text_frame.text == "Additional Slide":
            run = shape.text_frame.paragraphs[0].runs[0]
            run.text = "Impact Statement"
            run.font.size = Pt(17)
            shape.left, shape.top = Emu(329100), Emu(300000)
            shape.width, shape.height = Emu(8485800), Emu(400000)
    add_content_textbox(impact_slide, Emu(329100), Emu(950000), Emu(8485800), Emu(3400000), [
        ("Measurable outcomes", True),
        ("Zero drift spread across 14 metrics / 28 bindings on the last run. The divergence is priced "
         "in decisions, not decimals: 12 suppliers (22% of the \u201ccompliant\u201d list) cleared without "
         "earning it, 0 false escalations. 60 conversational questions scored end-to-end \u2014 the "
         "evaluation itself found and fixed a real broken access control (Q53). 162/170 automated "
         "tests passing.", False),
        ("Scalability potential", True),
        ("The pattern \u2014 one definition, one binding per consumer, one independent drift test \u2014 is "
         "domain-agnostic. Adding a metric is a registration, not a new pipeline; the same governance "
         "layer covers finance, quality, or any team that argues over whose number is right.", False),
        ("How this extends beyond the demo", True),
        ("The drift test already runs daily at 06:00 UTC against production semantic views. Personas "
         "map onto real Snowflake roles and row access policies, not app-level mocks. The "
         "conversational layer refuses out-of-ontology questions instead of guessing \u2014 which is what "
         "makes it safe to hand to a real user, not just a judge.", False),
    ])

    prs.save(PATH)
    print("saved", PATH, "-", len(prs.slides.__iter__.__self__._sldIdLst), "slides")


if __name__ == "__main__":
    build()
